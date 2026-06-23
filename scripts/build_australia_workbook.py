#!/usr/bin/env python3
import argparse
import datetime as dt
import json
import os
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "dist"
AUD_CNY = float(os.environ.get("FX_AUD_CNY", "4.65"))
OUTPUT_TIMEZONE = os.environ.get("OUTPUT_TIMEZONE", "Asia/Shanghai")

BLUE = "1F4E78"
GREY = "808080"
LIGHT_BLUE = "DDEBF7"
YELLOW = "FFFF00"
WHITE = "FFFFFF"
THIN_BORDER = Border(
    left=Side(style="thin", color="9EADCC"),
    right=Side(style="thin", color="9EADCC"),
    top=Side(style="thin", color="9EADCC"),
    bottom=Side(style="thin", color="9EADCC"),
)

CATEGORY_ORDER = ["牛", "猪", "鸡", "鸡蛋", "生鲜乳", "饲料"]
EXCLUDED_INDICATORS = {"猪价周变化(c/kg)"}
SHEET_SPECS = {
    "牛": (
        "牛 — 价格、成交头数、屠宰量与产量（月度统一）",
        "价格为月均；屠宰/产量为季度月份值；主指标为 EYCI 月均。",
        "EYCI月均(c/kg cwt)",
        "cents_to_rmb_kg",
    ),
    "猪": (
        "猪 — 行业价格、CPI、屠宰量与产量（月度统一）",
        "Australian Pork 周报取月均；屠宰/产量为季度月份值；主指标为 Buyers NAT 60.1-75kg 平均价。",
        "猪价 Buyers NAT 60.1-75kg平均(c/kg HSCW)",
        "cents_to_rmb_kg",
    ),
    "鸡": (
        "鸡 — CPI、屠宰量与产量（月度统一）",
        "鸡肉缺少公开生产端周度价，暂以 ABS Poultry CPI 表示零售价格趋势；主指标为禽肉 CPI。",
        "禽肉CPI指数",
        None,
    ),
    "鸡蛋": (
        "鸡蛋 — ABS CPI（月度）",
        "鸡蛋缺少公开生产端周度价，暂以 ABS Eggs CPI 表示零售价格趋势。",
        "鸡蛋CPI指数",
        None,
    ),
    "生鲜乳": (
        "生鲜乳 — ABS Milk CPI（月度）",
        "当前先用 ABS Milk CPI 表示零售价格趋势；农场价工具可后续单独接入。",
        "牛奶CPI指数",
        None,
    ),
    "饲料": (
        "饲料 — 澳洲本地饲料粮与蛋白粕价格（月度）",
        "Australian Pork Eyes & Ears 周报中的 ProFarmer 澳洲 Delivered 周价；每期按地区 TW 报价取均值，再按月平均。",
        "饲料小麦澳洲交付价月均(A$/t)",
        "aud_t_to_rmb_t",
    ),
}


def resolve_path(value):
    path = Path(value)
    if not path.is_absolute():
        path = ROOT / path
    return path


def find_latest_australia_csv():
    candidates = list((ROOT / "data").glob("澳洲农业数据_*_合并长表.csv"))
    if not candidates:
        raise SystemExit("找不到澳洲合并长表 CSV: data/澳洲农业数据_*_合并长表.csv")
    return max(candidates, key=lambda path: (path.stat().st_mtime, path.name))


def style_header(cell):
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(name="Noto Sans CJK SC", size=9, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def style_sheet(ws, max_row: int, max_col: int):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "D5"
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 7
    ws.column_dimensions["C"].width = 6
    for col in range(4, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 17
    for row in range(1, 5):
        ws.row_dimensions[row].height = 22 if row != 4 else 38
    for row in range(5, max_row + 1):
        fill = row % 2 == 0
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            cell.font = Font(name="Arial", size=9)
            cell.border = THIN_BORDER
            if fill:
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row, 1).number_format = "yyyy-mm"


def wide_by_category(data: pd.DataFrame):
    result = {}
    filtered = data[~data["指标"].isin(EXCLUDED_INDICATORS)].copy()
    filtered["日期"] = pd.to_datetime(filtered["日期"])
    filtered["数值"] = pd.to_numeric(filtered["数值"], errors="coerce")
    for category in CATEGORY_ORDER:
        part = filtered[filtered["品类"] == category]
        if part.empty:
            continue
        wide = (
            part.pivot_table(index="日期", columns="指标", values="数值", aggfunc="last")
            .sort_index()
            .reset_index()
        )
        wide.insert(1, "年", wide["日期"].dt.year)
        wide.insert(2, "月", wide["日期"].dt.month)
        result[category] = wide
    return result


def data_range_text(df: pd.DataFrame, main_col: str):
    if df.empty or main_col not in df.columns:
        return "无可用数据"
    usable = df[df[main_col].notna()]
    if usable.empty:
        return "主指标暂无数据"
    return f"{usable['日期'].min().date()} ~ {usable['日期'].max().date()} | 主: {main_col}"


def write_info_sheet(wb: Workbook, summaries: list[dict], export_date: str, source_file: Path):
    ws = wb.active
    ws.title = "参数与说明"
    rows = [
        ["澳洲农业数据 — 按品类月度口径", None, None],
        ["数据来源：MLA/NLRS、Australian Pork/ProFarmer、ABS Livestock Products、ABS CPI", None, None],
        [f"导出日期：{export_date}；来源文件：{source_file.name}", None, None],
        [None, None, None],
        ["可调参数", None, None],
        ["澳元兑人民币汇率", AUD_CNY, "用于澳洲 c/kg 与 A$/t 的统一人民币口径换算。"],
        [None, None, None],
        ["时间口径", "月度", "日度/周度价格取月均；ABS 季度屠宰/产量只在季度月份填值；CPI 为月度指数。"],
        ["口径提醒", None, "牛/猪/饲料价格多为生产端或行业报价；鸡/鸡蛋/牛奶当前使用 ABS CPI 零售价格指数，不与生产端价格直接比较。"],
        [None, None, None],
        ["各Sheet概览", None, None],
        ["Sheet", "行数", "时间范围 / 主指标"],
    ]
    for summary in summaries:
        rows.append([summary["sheet"], summary["rows"], summary["range"]])
    for row in rows:
        ws.append(row)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 100
    ws["A1"].font = Font(name="Noto Sans CJK SC", size=16, bold=True, color=BLUE)
    ws["B6"].fill = PatternFill("solid", fgColor=YELLOW)
    ws["B6"].font = Font(name="Arial", size=11, bold=True, color="0000FF")
    for cell in ws[11]:
        style_header(cell)


def write_category_sheet(wb: Workbook, name: str, df: pd.DataFrame):
    title, note, main_col, convert_kind = SHEET_SPECS[name]
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A2"] = note
    ws["A1"].font = Font(name="Noto Sans CJK SC", size=13, bold=True, color=BLUE)
    ws["A2"].font = Font(name="Noto Sans CJK SC", size=9, color=GREY)

    headers = list(df.columns) + ["主指标(统一口径)", "环比%", "同比%"]
    for col, header in enumerate(headers, 1):
        style_header(ws.cell(4, col, "月份" if header == "日期" else header))

    main_idx = list(df.columns).index(main_col) + 1 if main_col in df.columns else None
    calc_col = len(df.columns) + 1
    mom_col = calc_col + 1
    yoy_col = calc_col + 2
    for idx, record in df.iterrows():
        row = idx + 5
        for col, column in enumerate(df.columns, 1):
            value = record[column]
            ws.cell(row, col, None if pd.isna(value) else value)
        if not main_idx:
            continue
        letter = get_column_letter(main_idx)
        if convert_kind == "cents_to_rmb_kg":
            ws.cell(row, calc_col, f'=IF({letter}{row}="","",{letter}{row}*\'参数与说明\'!$B$6/100)')
            ws.cell(row, calc_col).number_format = "#,##0.00"
        elif convert_kind == "aud_t_to_rmb_t":
            ws.cell(row, calc_col, f'=IF({letter}{row}="","",{letter}{row}*\'参数与说明\'!$B$6)')
            ws.cell(row, calc_col).number_format = "#,##0"
        else:
            ws.cell(row, calc_col, f'=IF({letter}{row}="","",{letter}{row})')
            ws.cell(row, calc_col).number_format = "#,##0.00"
        ws.cell(row, mom_col, f'=IFERROR({letter}{row}/{letter}{row-1}-1,"")' if row > 5 else "")
        ws.cell(row, yoy_col, f'=IFERROR({letter}{row}/{letter}{row-12}-1,"")' if row > 16 else "")
        ws.cell(row, mom_col).number_format = "0.0%"
        ws.cell(row, yoy_col).number_format = "0.0%"

    max_row = len(df) + 4
    max_col = len(headers)
    style_sheet(ws, max_row, max_col)
    for row in range(5, max_row + 1):
        for col in range(4, len(df.columns) + 1):
            ws.cell(row, col).number_format = "#,##0.00"
    return {"sheet": name, "rows": len(df), "range": data_range_text(df, main_col)}


def parse_args():
    parser = argparse.ArgumentParser(description="从澳洲农业数据合并长表生成按品类 Excel")
    parser.add_argument("--input", help="澳洲农业数据合并长表 CSV；不传则自动找 data/ 下最新文件")
    parser.add_argument("--output-dir", default=str(OUT), help="输出目录")
    parser.add_argument("--output", help="指定输出 xlsx 路径")
    return parser.parse_args()


def main():
    args = parse_args()
    input_csv = resolve_path(args.input) if args.input else find_latest_australia_csv()
    output_dir = resolve_path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    now = dt.datetime.now(dt.UTC).astimezone(ZoneInfo(OUTPUT_TIMEZONE))
    stamp = now.strftime("%Y-%m-%d_%H%M")
    export_date = now.strftime("%Y-%m-%d")
    output = resolve_path(args.output) if args.output else output_dir / f"澳洲农业数据_{stamp}.xlsx"

    data = pd.read_csv(input_csv, encoding="utf-8-sig")
    sheets = wide_by_category(data)
    wb = Workbook()
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except AttributeError:
        pass
    summaries = []
    write_info_sheet(wb, summaries, export_date, input_csv)
    for category in CATEGORY_ORDER:
        if category not in sheets:
            continue
        summaries.append(write_category_sheet(wb, category, sheets[category]))
    wb.remove(wb["参数与说明"])
    wb.create_sheet("参数与说明", 0)
    wb.active = 0
    write_info_sheet(wb, summaries, export_date, input_csv)
    wb.save(output)
    print(json.dumps({"output": str(output), "input": str(input_csv), "sheets": len(sheets)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
