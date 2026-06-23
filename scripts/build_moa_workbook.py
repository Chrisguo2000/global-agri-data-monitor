#!/usr/bin/env python3
import argparse
import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


LONG_DATE_FORMAT = 'yyyy"年"m"月"d"日"'
DATE_AXIS_FORMAT = 'm"月"d"日"yy"年"'
CHART_WIDTH = 15
CHART_HEIGHT = 7.5
LINE_WIDTH = 18000
PALETTE = ["D93530", "1F2C4A", "A07A4B", "2CB56C", "5B7CFA", "D88425", "7F59B0"]

MARKET_ITEMS = sorted(
    [
        "育肥猪配合饲料",
        "肉鸡配合饲料",
        "蛋鸡配合饲料",
        "商品代蛋雏鸡",
        "商品代肉雏鸡",
        "商品代雏鸡",
        "生鲜乳",
        "仔猪",
        "生猪",
        "猪肉",
        "鸡蛋",
        "鸡肉",
        "牛肉",
        "活牛",
        "羊肉",
        "活羊",
        "玉米",
        "豆粕",
    ],
    key=len,
    reverse=True,
)


def compact(text=""):
    return (
        str(text or "")
        .replace("\u00a0", " ")
        .replace("\n", " ")
        .replace("／", "/")
        .replace("—", "-")
        .replace("－", "-")
    )


def normalize(text=""):
    text = compact(text)
    text = re.sub(r"\s*([，。；：、（）%])\s*", r"\1", text)
    text = re.sub(r"\s*/\s*", "/", text)
    text = re.sub(r"(\d)\s+(\d)", r"\1\2", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def parse_date(value):
    if not value:
        return None
    if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
        return value
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    return None


def pct(direction, raw):
    if raw is None or raw == "":
        return None
    value = float(raw)
    if direction in ("下降", "下跌"):
        return -abs(value)
    if direction in ("上涨", "增长"):
        return abs(value)
    if direction == "持平":
        return 0
    return value


def parse_change(tail, mode):
    labels = r"(?:环比|较前一周|比前一周|比上周|较上周)" if mode == "mom" else r"(?:同比|较去年同期)"
    if re.search(labels + r"[^。；，,]*持平", tail):
        return "持平", 0
    match = re.search(labels + r"[^。；，,]*(上涨|增长|下降|下跌)\s*([0-9]+(?:\.[0-9]+)?)\s*%", tail)
    if not match:
        return "", None
    return match.group(1), pct(match.group(1), match.group(2))


def split_sentences(text):
    text = normalize(text)
    return [s.strip() for s in re.split(r"(?<=[。；;])", text) if s.strip()]


def split_market_prefix(prefix):
    prefix = normalize(prefix).replace("全国 ", "全国")
    for item in MARKET_ITEMS:
        idx = prefix.rfind(item)
        if idx >= 0:
            scope = prefix[:idx].strip(" ，,。；;")
            if not scope:
                scope = "全国"
            return scope, item
    return "", ""


def title_period(article):
    match = re.search(r"(\d+月第\d+周)", article.get("title", ""))
    return match.group(1) if match else ""


def collection_date(article):
    match = re.search(r"采集日为\s*(\d+月\d+日)", compact(article.get("text", "")))
    return match.group(1) if match else ""


def parse_market_article(article):
    rows = []
    if article.get("type") != "畜产品和饲料集贸市场价格":
        return rows
    current_category = ""
    for sentence in split_sentences(article.get("text", "")):
        category_only = re.match(r"^([^。；;]{2,16}价格)[。；;]$", sentence)
        if category_only:
            current_category = category_only.group(1)
            continue
        category_with_text = re.match(r"^([^。；;]{2,16}价格)[。；;](.+)$", sentence)
        if category_with_text:
            current_category = category_with_text.group(1)
            sentence = category_with_text.group(2)

        price_re = re.compile(
            r"([^。；;，,]{0,90}?)(?:平均\s*价格|平均价格|价格)\s*(?:为)?\s*"
            r"([0-9]+(?:\.[0-9]+)?)\s*元/(公斤|只|吨|千克)([^。；;]*)"
        )
        for match in price_re.finditer(sentence):
            scope, item = split_market_prefix(match.group(1))
            if not item:
                continue
            tail = match.group(4) or ""
            mom_dir, mom_pct = parse_change(tail.replace("与前一周", "比前一周"), "mom")
            yoy_dir, yoy_pct = parse_change(tail, "yoy")
            rows.append(
                {
                    "publishDate": article.get("publishDate", ""),
                    "period": title_period(article),
                    "collectionDate": collection_date(article),
                    "category": current_category,
                    "scope": scope,
                    "item": item,
                    "metric": "平均价格" if "平均" in match.group(0) else "价格",
                    "value": float(match.group(2)),
                    "unit": f"元/{match.group(3)}",
                    "momDirection": mom_dir,
                    "momPct": mom_pct,
                    "yoyDirection": yoy_dir,
                    "yoyPct": yoy_pct,
                    "title": article.get("title", ""),
                    "url": article.get("url", ""),
                    "context": sentence,
                }
            )
    return dedupe(rows, ["publishDate", "scope", "item", "value", "context"])


def parse_slaughter_article(article):
    rows = []
    if article.get("type") != "生猪定点屠宰价格":
        return rows
    text = normalize(article.get("text", ""))
    monitored = re.search(r"据农业农村部监测，\s*([^，。；]*?日)\s*[，,]", text)
    specs = [
        ("生猪平均收购价格", r"生猪\s*平均收购价格\s*为?\s*([0-9]+(?:\.[0-9]+)?)\s*元/公斤([^。；;]*)"),
        ("白条肉平均出厂价格", r"白条肉\s*平均出厂价格\s*为?\s*([0-9]+(?:\.[0-9]+)?)\s*元/公斤([^。；;]*)"),
    ]
    for item, pattern in specs:
        for match in re.finditer(pattern, text):
            tail = match.group(2) or ""
            mom_dir, mom_pct = parse_change(tail, "mom")
            yoy_dir, yoy_pct = parse_change(tail, "yoy")
            rows.append(
                {
                    "publishDate": article.get("publishDate", ""),
                    "period": title_period(article),
                    "monitoredPeriod": monitored.group(1) if monitored else "",
                    "scope": "全国",
                    "item": item,
                    "metric": "平均价格",
                    "value": float(match.group(1)),
                    "unit": "元/公斤",
                    "momDirection": mom_dir,
                    "momPct": mom_pct,
                    "yoyDirection": yoy_dir,
                    "yoyPct": yoy_pct,
                    "title": article.get("title", ""),
                    "url": article.get("url", ""),
                    "context": text,
                }
            )
    return dedupe(rows, ["publishDate", "item", "value", "url"])


def dedupe(rows, keys):
    seen = set()
    result = []
    for row in rows:
        key = tuple(row.get(k, "") for k in keys)
        if key in seen:
            continue
        seen.add(key)
        result.append(row)
    return result


def clear_sheet_rows(ws, header_row):
    if ws.max_row > header_row:
        ws.delete_rows(header_row + 1, ws.max_row - header_row)


def header_map(ws, header_row):
    return {str(ws.cell(header_row, col).value or "").strip(): col for col in range(1, ws.max_column + 1)}


def scope_all(scope):
    return scope == "全国"


def scope_non_all(scope):
    return scope != "全国"


def scope_contains(text):
    return lambda scope: text in scope


def find_metric(rows, item, scope_pred):
    for row in rows:
        if row["item"] == item and scope_pred(row["scope"]):
            return row
    return None


def set_metric(out, headers, rows, price_header, item, scope_pred, mom_header=None, yoy_header=None):
    metric = find_metric(rows, item, scope_pred)
    if not metric:
        return
    if price_header in headers:
        out[headers[price_header]] = metric["value"]
    if mom_header and mom_header in headers:
        out[headers[mom_header]] = metric["momPct"]
    if yoy_header and yoy_header in headers:
        out[headers[yoy_header]] = metric["yoyPct"]


SHEET_CONFIGS = {
    "生猪产品价格": [
        ("仔猪价格（元/kg）", "仔猪", scope_all, "仔猪环比", "仔猪同比"),
        ("生猪价格（元/kg）", "生猪", scope_all, "生猪环比", "生猪同比"),
        ("猪肉价格（元/kg）", "猪肉", scope_all, "猪肉环比", "猪肉同比"),
    ],
    "家禽产品价格": [
        ("鸡蛋价格（元/kg）", "鸡蛋", scope_all, "鸡蛋环比", "鸡蛋同比"),
        ("鸡蛋(主产省份)价格（元/kg）", "鸡蛋", scope_non_all, "鸡蛋(主产省份)环比", "鸡蛋(主产省份)同比"),
        ("鸡肉价格（元/kg）", "鸡肉", scope_all, "鸡肉环比", "鸡肉同比"),
        ("商品代蛋雏鸡价格（元/只）", "商品代蛋雏鸡", scope_all, "商品代蛋雏鸡环比", "商品代蛋雏鸡同比"),
        ("商品代肉雏鸡价格（元/只）", "商品代肉雏鸡", scope_all, "商品代肉雏鸡环比", "商品代肉雏鸡同比"),
    ],
    "牛羊产品价格": [
        ("牛肉价格（元/kg）", "牛肉", scope_all, "牛肉环比", "牛肉同比"),
        ("牛肉(主产省份)价格（元/kg）", "牛肉", scope_non_all, "牛肉(主产省份)环比", "牛肉(主产省份)同比"),
        ("活牛(主产省份)价格（元/kg）", "活牛", scope_non_all, "活牛(主产省份)环比", "活牛(主产省份)同比"),
        ("羊肉价格（元/kg）", "羊肉", scope_all, "羊肉环比", "羊肉同比"),
        ("羊肉(主产省份)价格（元/kg）", "羊肉", scope_non_all, "羊肉(主产省份)环比", "羊肉(主产省份)同比"),
        ("活羊(主产省份)价格（元/kg）", "活羊", scope_non_all, "活羊(主产省份)环比", "活羊(主产省份)同比"),
    ],
    "生鲜乳价格": [
        ("生鲜乳(主产省份)价格（元/kg）", "生鲜乳", lambda scope: True, "生鲜乳(主产省份)环比", "生鲜乳(主产省份)同比"),
    ],
    "饲料价格": [
        ("玉米价格（元/kg）", "玉米", scope_all, "玉米环比", "玉米同比"),
        ("玉米(东北三省)价格（元/kg）", "玉米", scope_contains("东北三省"), "玉米(东北三省)环比", "玉米(东北三省)同比"),
        ("玉米(广东省)价格（元/kg）", "玉米", scope_contains("广东"), "玉米(广东省)环比", "玉米(广东省)同比"),
        ("豆粕价格（元/kg）", "豆粕", scope_all, "豆粕环比", "豆粕同比"),
        ("育肥猪配合饲料价格（元/kg）", "育肥猪配合饲料", scope_all, "育肥猪配合饲料环比", "育肥猪配合饲料同比"),
        ("肉鸡配合饲料价格（元/kg）", "肉鸡配合饲料", scope_all, "肉鸡配合饲料环比", "肉鸡配合饲料同比"),
        ("蛋鸡配合饲料价格（元/kg）", "蛋鸡配合饲料", scope_all, "蛋鸡配合饲料环比", "蛋鸡配合饲料同比"),
    ],
}


def sorted_dates(rows):
    return sorted({row["publishDate"] for row in rows if row.get("publishDate")}, reverse=True)


def date_key(value):
    parsed = parse_date(value)
    if parsed:
        return parsed.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    return text


def read_existing_summary(ws, header_row):
    existing = {}
    for row in range(header_row + 1, ws.max_row + 1):
        key = date_key(ws.cell(row, 1).value)
        if not key:
            continue
        existing[key] = {col: ws.cell(row, col).value for col in range(1, ws.max_column + 1)}
    return existing


def sorted_date_keys(keys):
    return sorted(keys, key=lambda value: parse_date(value) or datetime.min.date(), reverse=True)


def populate_market_summary(ws, rows):
    header_row = 1 if ws.title == "生猪产品价格" else 2
    existing = read_existing_summary(ws, header_row)
    headers = header_map(ws, header_row)
    by_date = defaultdict(list)
    for row in rows:
        by_date[row["publishDate"]].append(row)
    clear_sheet_rows(ws, header_row)
    all_dates = sorted_date_keys(set(existing) | set(sorted_dates(rows)))
    for row_idx, publish_date in enumerate(all_dates, start=header_row + 1):
        current = by_date.get(publish_date, [])
        out = dict(existing.get(publish_date, {}))
        out[1] = parse_date(publish_date)
        if current:
            out[2] = current[0].get("period", "")
        for price_header, item, scope_pred, mom_header, yoy_header in SHEET_CONFIGS[ws.title]:
            set_metric(out, headers, current, price_header, item, scope_pred, mom_header, yoy_header)
        for col, value in out.items():
            cell = ws.cell(row_idx, col)
            cell.value = value
            if col == 1:
                cell.number_format = LONG_DATE_FORMAT
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 16)


def populate_slaughter_summary(ws, rows):
    header_row = 2
    existing = read_existing_summary(ws, header_row)
    headers = header_map(ws, header_row)
    by_date = defaultdict(list)
    for row in rows:
        by_date[row["publishDate"]].append(row)
    clear_sheet_rows(ws, header_row)
    all_dates = sorted_date_keys(set(existing) | set(sorted_dates(rows)))
    for row_idx, publish_date in enumerate(all_dates, start=header_row + 1):
        current = by_date.get(publish_date, [])
        out = dict(existing.get(publish_date, {}))
        out[1] = parse_date(publish_date)
        if current:
            out[2] = current[0].get("period", "")
        set_metric(out, headers, current, "生猪平均收购价格（元/kg）", "生猪平均收购价格", scope_all, "生猪平均收购价格环比", "生猪平均收购价格同比")
        set_metric(out, headers, current, "白条肉平均出厂价格（元/kg）", "白条肉平均出厂价格", scope_all, "白条肉平均出厂价格环比", "白条肉平均出厂价格同比")
        for col, value in out.items():
            cell = ws.cell(row_idx, col)
            cell.value = value
            if col == 1:
                cell.number_format = LONG_DATE_FORMAT
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 16)


def is_price_header(header):
    text = str(header or "").strip()
    return "价格（" in text and "环比" not in text and "同比" not in text


def add_chart(ws):
    header_row = 1 if ws.title == "生猪产品价格" else 2
    first_data_row = header_row + 1
    last_data_row = ws.max_row
    price_cols = [col for col in range(1, ws.max_column + 1) if is_price_header(ws.cell(header_row, col).value)]
    if not price_cols or last_data_row < first_data_row:
        return None
    ws._charts = []
    chart = LineChart()
    chart.title = "" if ws.title == "生鲜乳价格" else f"{ws.title} 全量价格走势"
    chart.width = CHART_WIDTH
    chart.height = CHART_HEIGHT
    if ws.title != "生鲜乳价格":
        chart.y_axis.title = "价格"
        chart.x_axis.title = "时间"
        chart.legend.position = "b"
    else:
        chart.legend = None
    chart.x_axis.number_format = DATE_AXIS_FORMAT
    chart.x_axis.tickLblPos = "nextTo"
    data = Reference(ws, min_col=min(price_cols), max_col=max(price_cols), min_row=header_row, max_row=last_data_row)
    categories = Reference(ws, min_col=1, min_row=first_data_row, max_row=last_data_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    for idx, series in enumerate(chart.series):
        color = PALETTE[idx % len(PALETTE)]
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 28575 if ws.title == "生鲜乳价格" else LINE_WIDTH
    anchor = f"{get_column_letter(ws.max_column + 2)}2"
    ws.add_chart(chart, anchor)
    return anchor


def read_existing_table(ws, columns):
    header_lookup = {str(ws.cell(1, col).value or "").strip(): col for col in range(1, ws.max_column + 1)}
    rows = []
    for row_idx in range(2, ws.max_row + 1):
        record = {}
        has_value = False
        for key, header, _width in columns:
            col = header_lookup.get(header)
            value = ws.cell(row_idx, col).value if col else ""
            if value not in ("", None):
                has_value = True
            if key == "publishDate":
                parsed = parse_date(value)
                record[key] = parsed.strftime("%Y-%m-%d") if parsed else value
            else:
                record[key] = value if value is not None else ""
        if has_value:
            rows.append(record)
    return rows


def merge_records(existing, incoming, keys):
    merged = {}
    order = []
    for record in existing + incoming:
        key = tuple(record.get(k, "") for k in keys)
        if key not in merged:
            order.append(key)
        merged[key] = record
    return [merged[key] for key in order]


def write_table(wb, name, columns, rows, merge_keys=None):
    ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
    if merge_keys:
        rows = merge_records(read_existing_table(ws, columns), rows, merge_keys)
        rows = sorted(rows, key=lambda row: (parse_date(row.get("publishDate")) or datetime.min.date(), str(row.get("ordinal", ""))), reverse=True)
    ws.delete_rows(1, ws.max_row)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for col_idx, (key, header, width) in enumerate(columns, start=1):
        cell = ws.cell(1, col_idx)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (key, _header, _width) in enumerate(columns, start=1):
            cell = ws.cell(row_idx, col_idx)
            value = row.get(key, "")
            if key == "publishDate":
                value = parse_date(value) or value
                cell.number_format = LONG_DATE_FORMAT
            cell.value = value
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def latest_publish_date(articles):
    dates = [a.get("publishDate", "") for a in articles if a.get("publishDate")]
    return max(dates) if dates else datetime.now().strftime("%Y-%m-%d")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    articles = data.get("articles", [])
    market_rows = []
    slaughter_rows = []
    for article in articles:
        market_rows.extend(parse_market_article(article))
        slaughter_rows.extend(parse_slaughter_article(article))

    wb = load_workbook(args.template)
    for name in SHEET_CONFIGS:
        populate_market_summary(wb[name], market_rows)
        add_chart(wb[name])
    populate_slaughter_summary(wb["生猪定点屠宰价格"], slaughter_rows)
    add_chart(wb["生猪定点屠宰价格"])

    write_table(
        wb,
        "文章目录",
        [
            ("ordinal", "序号", 8),
            ("listPage", "列表页", 8),
            ("publishDate", "发布日期", 16),
            ("type", "类型", 22),
            ("title", "标题", 55),
            ("source", "来源", 24),
            ("url", "URL", 70),
        ],
        articles,
        merge_keys=["url"],
    )
    write_table(
        wb,
        "集贸价格明细",
        [
            ("publishDate", "发布日期", 16),
            ("period", "标题周期", 12),
            ("collectionDate", "采集日", 12),
            ("category", "分类", 16),
            ("scope", "范围", 28),
            ("item", "指标项", 20),
            ("value", "数值", 10),
            ("unit", "单位", 10),
            ("momPct", "环比(%)", 10),
            ("yoyPct", "同比(%)", 10),
            ("context", "原句", 90),
            ("url", "URL", 70),
        ],
        market_rows,
        merge_keys=["url", "scope", "item", "value"],
    )
    write_table(
        wb,
        "屠宰价格明细",
        [
            ("publishDate", "发布日期", 16),
            ("period", "标题周期", 12),
            ("monitoredPeriod", "监测区间", 18),
            ("item", "指标项", 24),
            ("value", "数值", 10),
            ("unit", "单位", 10),
            ("momPct", "环比(%)", 10),
            ("yoyPct", "同比(%)", 10),
            ("context", "原句", 90),
            ("url", "URL", 70),
        ],
        slaughter_rows,
        merge_keys=["url", "item", "value"],
    )

    try:
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except AttributeError:
        pass

    output = Path(args.output)
    if output.suffix.lower() != ".xlsx":
        date_tag = latest_publish_date(articles).replace("-", "")
        output.mkdir(parents=True, exist_ok=True)
        output = output / f"农业农村部监测预警数据汇总_{date_tag}_清洗版.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(json.dumps({"output": str(output), "marketRows": len(market_rows), "slaughterRows": len(slaughter_rows)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
